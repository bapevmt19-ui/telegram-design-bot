"""Quản lý tài chính: /salary /budget /spend /goal /report /report_excel"""
import asyncio
import logging
import os
from datetime import datetime

import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.constants import ParseMode
from telegram.ext import ContextTypes

from ai_client import clean_for_telegram, parse_amount
from core_actions import execute_spend
from storage import finance_store
from telegram_helpers import send_chunked_message, unique_temp_path

logger = logging.getLogger(__name__)


async def salary_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        amount = parse_amount(context.args[0])

        def _mutate(data):
            data["salary"] = amount
            data["timo_confirmed"] = False
            return data

        await finance_store.update(_mutate)
        msg = (
            f"💰 <b>ĐÃ GHI NHẬN LƯƠNG:</b> {amount:,.0f} VNĐ\n"
            "Sếp hãy chia /budget và bấm nút dưới đây để xác nhận đã thao tác trên app Timo!"
        )
        keyboard = [[InlineKeyboardButton("🏦 Đã chia tiền vào các hũ Timo", callback_data="timo_confirm")]]
        await update.message.reply_text(
            clean_for_telegram(msg), parse_mode=ParseMode.HTML, reply_markup=InlineKeyboardMarkup(keyboard)
        )
    except Exception:
        await update.message.reply_text("Lỗi cú pháp! Gõ: /salary [số tiền]")


async def budget_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        category = context.args[0].replace("_", " ").title()
        amount = parse_amount(context.args[1])

        def _mutate(data):
            data.setdefault("budgets", {})[category] = amount
            return data

        await finance_store.update(_mutate)
        await send_chunked_message(
            update.message.reply_text, f"🎯 <b>CẬP NHẬT QUỸ: {category}</b>\n• Hạn mức: {amount:,.0f} VNĐ"
        )
    except Exception:
        await update.message.reply_text("Lỗi cú pháp! Gõ: /budget [tên hũ] [số tiền]")


async def spend_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        amount = parse_amount(context.args[0])
        reason = " ".join(context.args[1:])
        await execute_spend(update.message.chat_id, context, amount, reason)
    except Exception:
        await update.message.reply_text("Lỗi cú pháp! Gõ: /spend [số tiền] [lý do]")


def _render_pie_chart(labels, sizes, title, path):
    """Hàm CPU-bound thuần (không async), được chạy trong thread
    riêng bởi asyncio.to_thread ở report_command.

    try/finally đảm bảo plt.close(fig) LUÔN được gọi kể cả khi
    ax.pie() lỗi (VD dữ liệu rỗng/âm) — bản gốc gọi plt.close() ở
    ngay sau dòng có thể raise, nên nếu lỗi xảy ra thì figure không
    bao giờ được đóng -> rò rỉ bộ nhớ matplotlib tích luỹ dần qua
    nhiều lần /report lỗi khi bot chạy dài ngày.
    """
    fig, ax = plt.subplots(figsize=(7, 7))
    try:
        ax.pie(sizes, labels=labels, autopct="%1.1f%%", startangle=90, colors=plt.cm.Paired.colors)
        ax.axis("equal")
        plt.title(title)
        plt.savefig(path, bbox_inches="tight")
    finally:
        plt.close(fig)


async def report_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Vẽ biểu đồ chi tiêu"""
    finance = await finance_store.read()
    current_month = datetime.now().strftime("%Y-%m")
    month_expenses = [item for item in finance.get("expenses", []) if item["date"].startswith(current_month)]

    if not month_expenses:
        await update.message.reply_text("Tháng này sếp chưa tiêu đồng nào cả!")
        return

    categories = {}
    for exp in month_expenses:
        reason_words = exp["reason"].split()
        # BUG FIX: bản gốc dùng exp['reason'].split()[0] làm fallback,
        # sẽ IndexError nếu reason rỗng.
        fallback = reason_words[0].title() if reason_words else "Khac"
        cat = exp.get("category", fallback)
        categories[cat] = categories.get(cat, 0) + exp["amount"]

    chart_path = unique_temp_path("chart", ".png")
    title = f"Phân bổ chi tiêu tháng {current_month}"
    # Vẽ chart là việc CPU-bound + ghi đĩa -> chạy trong thread riêng
    # để không chặn event loop của bot trong lúc render.
    await asyncio.to_thread(
        _render_pie_chart, list(categories.keys()), list(categories.values()), title, chart_path
    )

    try:
        with open(chart_path, "rb") as photo:
            await update.message.reply_photo(
                photo=photo, caption=f"📊 Báo cáo phân bổ chi tiêu tháng {current_month}"
            )
    finally:
        if os.path.exists(chart_path):
            os.remove(chart_path)


def _render_excel_report(month_expenses: list, current_month: str, path: str):
    """Hàm CPU/IO-bound thuần (không async), chạy trong thread riêng
    bởi asyncio.to_thread ở report_excel_command — cùng nguyên tắc với
    _render_pie_chart phía trên (không chặn event loop của bot).

    Nâng cấp (17/9, lần 7): xuất báo cáo chi tiêu tháng ra file Excel
    dạng bảng có cột được canh chỉnh (Ngày / Số tiền / Danh mục / Lý
    do), có dòng tổng cuối bảng — bổ sung bên cạnh /report (biểu đồ
    tròn) chứ không thay thế, vì có nhu cầu xem số liệu dạng bảng để
    dễ đối chiếu/lọc hơn là chỉ nhìn biểu đồ.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"Chi tieu {current_month}"[:31]  # Excel giới hạn tên sheet 31 ký tự

    headers = ["Ngày", "Số tiền (VNĐ)", "Danh mục", "Lý do"]
    ws.append(headers)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total = 0
    for exp in sorted(month_expenses, key=lambda e: e["date"]):
        ws.append([exp["date"], exp["amount"], exp.get("category", "Khac"), exp["reason"]])
        total += exp["amount"]

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=1, value="TỔNG CỘNG").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=2, value=total)
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0"

    for row in ws.iter_rows(min_row=2, max_row=total_row, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "#,##0"

    for col_idx, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))]
            + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    wb.save(path)


async def report_excel_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    finance = await finance_store.read()
    current_month = datetime.now().strftime("%Y-%m")
    month_expenses = [e for e in finance.get("expenses", []) if e["date"].startswith(current_month)]

    if not month_expenses:
        await update.message.reply_text("Tháng này sếp chưa tiêu đồng nào cả!")
        return

    excel_path = unique_temp_path("report", ".xlsx")
    await asyncio.to_thread(_render_excel_report, month_expenses, current_month, excel_path)

    try:
        with open(excel_path, "rb") as f:
            await update.message.reply_document(
                document=f,
                filename=f"bao_cao_chi_tieu_{current_month}.xlsx",
                caption=f"📊 Báo cáo chi tiêu tháng {current_month} dạng bảng.",
            )
    finally:
        if os.path.exists(excel_path):
            os.remove(excel_path)


async def goal_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        item_name = " ".join(context.args[:-1])
        goal_amount = parse_amount(context.args[-1])
        finance = await finance_store.read()
        salary = finance.get("salary", 0)
        total_budget = sum(finance.get("budgets", {}).values())
        monthly_saving = salary - total_budget

        if monthly_saving <= 0:
            msg = (
                f"⚠️ Sếp không còn tiền dư mỗi tháng (Thu: {salary:,.0f}, "
                f"Chi: {total_budget:,.0f}). Không thể tiết kiệm!"
            )
        else:
            months_needed = goal_amount / monthly_saving
            msg = f"🎯 <b>MỤC TIÊU: {item_name.upper()}</b>\n• Thời gian dự kiến: <b>{months_needed:.1f} tháng</b>"
        await send_chunked_message(update.message.reply_text, clean_for_telegram(msg))
    except Exception:
        # BUG FIX: bản gốc "except: pass" -> sếp gõ sai cú pháp thì
        # không nhận được phản hồi gì cả, tưởng bot bị treo.
        await update.message.reply_text("Lỗi cú pháp! Gõ: /goal [tên mục tiêu] [số tiền]")
